"""Build the editable workshop deck. Run: python slides/build_slides.py"""
from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = Path(__file__).resolve().parent
NOTES = json.loads((HERE / 'speaker_notes.json').read_text(encoding='utf-8'))
# Replace these before the event. QR and screenshot slots remain editable.
SHORT_LINK = '[SHORT_LINK]'
MODEL_ID = '[MODEL_ID]'
LINKEDIN = '[LINKEDIN]'
GITHUB = '[GITHUB]'
BG, PANEL, WHITE, GREY, GREEN, RED = ('0B0B0F','17171E','FFFFFF','C8CBD3','22C55E','EF4444')
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

def rgb(v): return RGBColor.from_string(v)
def text(s, value, x, y, w, h, size=28, color=GREY, bold=False, mono=False, center=False):
    shape = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.02)
    tf.margin_top = tf.margin_bottom = Inches(.02)
    for i,line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = 'Consolas' if mono else 'Aptos'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = rgb(color)
        p.space_after = Pt(10)
        if center: p.alignment=PP_ALIGN.CENTER
    return shape

def box(s, value, x,y,w,h, color=GREEN, size=28):
    shape=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb=rgb(PANEL)
    shape.line.color.rgb=rgb(color); shape.line.width=Pt(1.5)
    shape.adjustments[0]=.12
    shape.text_frame.clear(); shape.text_frame.word_wrap=True
    shape.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
    for i,line in enumerate(value.split('\n')):
        p=shape.text_frame.paragraphs[0] if i==0 else shape.text_frame.add_paragraph()
        p.text=line; p.alignment=PP_ALIGN.CENTER
        p.font.name='Aptos';p.font.size=Pt(size);p.font.color.rgb=rgb(WHITE)
    return shape

def arrow(s,x,y,w=.65,h=.27,direction='right'):
    typ={'right':MSO_SHAPE.RIGHT_ARROW,'left':MSO_SHAPE.LEFT_ARROW,'down':MSO_SHAPE.DOWN_ARROW,'up':MSO_SHAPE.UP_ARROW}[direction]
    a=s.shapes.add_shape(typ,Inches(x),Inches(y),Inches(w),Inches(h))
    a.fill.solid();a.fill.fore_color.rgb=rgb(GREEN);a.line.fill.background()

def slide(title=None):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
    if title: text(s,title,.7,.55,11.95,1.3,42,WHITE,True)
    s.notes_slide.notes_text_frame.text=NOTES[len(prs.slides)-1]
    return s

s=slide()
text(s,'Build Your First\nAI Agent',.8,1.0,11.8,2.3,60,WHITE,True)
text(s,'CodeDoctor 🩺',.85,3.65,11,1,40,GREEN)
text(s,'Raghav',.85,5.8,9,.6,28)

s=slide()
text(s,'A chatbot answers.',.8,1.8,11.8,1,50,WHITE,True)
text(s,'An agent acts.',.8,3.25,11.8,1.2,64,GREEN,True)

s=slide('The brain in a jar')
for x,label in [(0.8,'LLM\nbrain'),(4.7,'Tools\nhands'),(8.6,'Loop\npersistence')]:box(s,label,x,2.6,3.7,2.0,size=32)

s=slide()
text(s,'Agent = LLM + Tools + Loop',.5,2.85,12.33,1.25,44,GREEN,True,center=True)

s=slide('The loop')
for label,x in [('Ask',.8),('Think',3.8),('Use a tool',6.8),('Read result',9.8)]:box(s,label,x,2.45,2.6,1.05)
for x in [3.45,6.45,9.45]:arrow(s,x,2.83,.3,.25)
# Read result returns to Think through a clearly separated lower track.
arrow(s,10.95,3.65,.25,.65,'down')
line=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(5.1),Inches(4.45),Inches(11.08),Inches(4.45))
line.line.color.rgb=rgb(GREEN);line.line.width=Pt(2)
arrow(s,4.98,3.65,.25,.8,'up')
box(s,'Answer',3.8,5.3,2.6,.9)
text(s,'The model decides the steps.',7,5.35,5.4,1,28)
# Exit from Think runs to the left of the return path.
line=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(4.15),Inches(3.5),Inches(4.15),Inches(5.05))
line.line.color.rgb=rgb(GREY);line.line.width=Pt(2)
arrow(s,4.025,4.85,.25,.35,'down')

s=slide('Buzzwords, decoded')
text(s,'Workflow',.8,2.1,3.0,.6,30,GREEN,True)
text(s,'You decide the steps',4.0,2.1,8.3,.7,32)
text(s,'Agent',.8,3.35,3.0,.6,30,GREEN,True)
text(s,'The model decides',4.0,3.35,8.3,.7,32)
text(s,'Multi-agent',.8,4.6,3.1,.7,30,GREEN,True)
text(s,'Agents delegate to agents',4.0,4.6,8.3,.8,32)

s=slide('Tonight: CodeDoctor 🩺')
for label,x,y in [('Run',.8,2.25),('Read error',4.9,2.25),('Fix',9,2.25),('Run tests',.8,4.5),('Write tests',6.9,4.5)]:box(s,label,x,y,3.5,1.0,color=RED if label=='Read error' else GREEN)
arrow(s,4.4,2.62,.4,.25);arrow(s,8.5,2.62,.4,.25)
arrow(s,10.6,3.4,.25,.8,'down');arrow(s,10.5,4.88,.8,.25,'left')
arrow(s,4.65,4.88,1.9,.25,'left');arrow(s,2.4,3.4,.25,.8,'up')
text(s,'Repeat until green ✓',.8,6.25,11.7,.7,30,GREEN,center=True)

s=slide('Setup')
box(s,'[QR CODE\nPLACEHOLDER]',.8,2.0,3.0,3.0,color=GREY,size=24)
text(s,SHORT_LINK,.8,5.3,3.6,.8,24,GREEN)
text(s,'1  Save a copy in Drive\n2  Paste your key\n3  Run cell 1',4.5,2.05,8,2.25,32,WHITE)
text(s,'Model ID: '+MODEL_ID,4.5,4.7,8,.8,28,GREEN)
text(s,'Green = done',4.5,6.0,3.7,.6,26,GREEN)
text(s,'Red = stuck',8.5,6.0,3.7,.6,26,RED)

s=slide('Stage 1: The brain in a jar')
text(s,'Confident ≠ Correct',.7,3.0,12,1.4,60,GREEN,True,center=True)

s=slide('Stage 2: Giving it hands')
for x,label in [(.8,'run_code'),(4.7,'save_file'),(8.6,'run_tests')]:box(s,label,x,2.85,3.7,1.5,size=30)

s=slide('Stage 3: This for-loop IS the agent')
panel=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(.7),Inches(2.15),Inches(11.95),Inches(3.9))
panel.fill.solid();panel.fill.fore_color.rgb=rgb(PANEL);panel.line.fill.background()
text(s,'for step in range(1, max_steps + 1):\n    ...\n    if not msg.tool_calls: return msg.content\n    ...\n    result = TOOLS[call.function.name](**args)',1.0,2.55,11.4,3.1,25,WHITE,mono=True)

s=slide('Code that runs ≠ Code that works')
text(s,'Crash bugs',.8,2.0,4.5,.7,32,RED,True)
text(s,'NameError\nZeroDivisionError',.8,3.0,4.7,2,28)
text(s,'Logic bugs',6.0,2.0,6.5,.7,32,RED,True)
text(s,'Class average: 184.75 / 100\nTopper = worst student\nExactly 70 gets F',6.0,3.0,6.6,2.4,27)
text(s,"Tests are the agent's proof.",.8,6.2,11.7,.7,32,GREEN,center=True)

s=slide('Stage 4: Ship it 🚀')
text(s,'One line:',.8,2.2,11.7,.7,30)
text(s,'demo.launch(share=True)',.8,3.25,11.7,1,40,GREEN,mono=True)
text(s,'Open it on your phone',.8,5.3,11.7,.9,36,WHITE)

s=slide('Same loop. Bigger tools.')
box(s,'[CODELENS SCREENSHOT PLACEHOLDER]',.8,1.95,11.7,3.75,color=GREY,size=26)
text(s,'CodeLens: understand any codebase instantly',.8,6.15,11.7,.8,28,GREEN)

s=slide('Where next')
text(s,'Frameworks wrap the loop\nMCP connects tools\nMulti-agent delegates work',.8,2.05,11.7,2.4,34,WHITE)
text(s,'You already know the core.\nEverything else is plumbing.',.8,5.25,11.7,1.45,34,GREEN,True)

s=slide('Thank you!')
text(s,LINKEDIN+'\n'+GITHUB,.8,2.6,6.5,1.8,30,GREEN)
box(s,'[FEEDBACK QR\nPLACEHOLDER]',8.5,2.2,3.5,3.5,color=GREY,size=24)

assert len(prs.slides)==16
target=HERE/'CodeDoctor_Slides.pptx'
prs.save(target)
print(target)
